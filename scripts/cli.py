#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""cli.py — DUT 论文格式改写的命令行入口。

    # 把任意 docx 改写成 DUT 格式
    python scripts/cli.py reformat 我的论文.docx -o 我的论文_DUT.docx --title "基于XX的研究"

    # 只做体检，不改文件
    python scripts/cli.py audit 我的论文.docx

    # 从零建一份骨架
    python scripts/cli.py template 论文骨架.docx --title "题目" --author 张三 --sid 12345678

    # 打包成可分发的 .skill
    python scripts/package_skill.py
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document  # noqa: E402

from dut_thesis import (  # noqa: E402
    add_blank, add_cover, add_heading, add_list_of_field, add_plain_heading,
    add_toc_field, add_section, audit, enable_update_fields, new_doc,
    print_report, reformat, set_header_pair, set_page_number_footer,
)


def _save(doc, path: str):
    enable_update_fields(doc)
    doc.save(path)
    print(f'已保存: {path}')


def cmd_reformat(args):
    doc = Document(args.input)
    stats = reformat(doc, title_cn=args.title, strict_refs=not args.keep_all_refs)
    print_report(stats, '改写统计')
    _save(doc, args.output)
    print()
    print_report(audit(Document(args.output)), '产物校验')
    return 0


def cmd_audit(args):
    info = audit(Document(args.input))
    print_report(info, '格式体检')
    return 0 if info['ok'] else 1


def cmd_template(args):
    doc = new_doc()
    # 第 1 节：封面（无页码）
    add_cover(doc, title_cn=args.title, title_en=args.title_en,
              fields=[('作 者 姓 名', args.author), ('学　　　号', args.sid),
                      ('指 导 教 师', args.advisor), ('学科、 专业', args.major),
                      ('答 辩 日 期', args.date)])

    # 第 2 节：前置部分（罗马页码）
    add_section(doc)
    set_page_number_footer(doc.sections[-1], fmt='upperRoman', start=1, style='dash')
    set_header_pair(doc.sections[-1], odd_text='大连理工大学硕士学位论文',
                    even_text=args.title)
    add_plain_heading(doc, '摘　　要')
    add_blank(doc)
    add_plain_heading(doc, '关键词：', size=12, cn='黑体', center=False, after=0)
    add_section(doc)
    add_plain_heading(doc, 'ABSTRACT', cn='Times New Roman')
    add_section(doc)
    add_plain_heading(doc, '目　　录')
    add_toc_field(doc)
    add_section(doc)
    add_plain_heading(doc, '图目录')
    add_list_of_field(doc, 'figure')
    add_section(doc)
    add_plain_heading(doc, '表目录')
    add_list_of_field(doc, 'table')

    # 第 3 节：正文（阿拉伯页码从 1 起）
    add_section(doc)
    set_page_number_footer(doc.sections[-1], fmt='decimal', start=1, style='dash')
    set_header_pair(doc.sections[-1], odd_text='大连理工大学硕士学位论文',
                    even_text=args.title)
    add_heading(doc, '1 绪论', 1)
    add_heading(doc, '1.1 研究背景与意义', 2)
    add_heading(doc, '1.1.1 研究背景', 3)

    for sec in doc.sections:
        set_header_pair(sec, odd_text='大连理工大学硕士学位论文', even_text=args.title)

    _save(doc, args.output)
    print_report(audit(Document(args.output)), '产物校验')
    return 0


def main(argv=None):
    ap = argparse.ArgumentParser(prog='dut-thesis', description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)

    p = sub.add_parser('reformat', help='把 docx 改写成 DUT 格式')
    p.add_argument('input')
    p.add_argument('-o', '--output', required=True)
    p.add_argument('--title', default=None, help='论文中文题目（用于偶数页页眉）')
    p.add_argument('--keep-all-refs', action='store_true',
                   help='保留未被引用的参考文献（默认删除）')
    p.set_defaults(func=cmd_reformat)

    p = sub.add_parser('audit', help='体检，不修改文件')
    p.add_argument('input')
    p.set_defaults(func=cmd_audit)

    p = sub.add_parser('template', help='从零生成符合 DUT 格式的论文骨架')
    p.add_argument('output')
    p.add_argument('--title', required=True)
    p.add_argument('--title-en', default='')
    p.add_argument('--author', default='')
    p.add_argument('--sid', default='')
    p.add_argument('--advisor', default='')
    p.add_argument('--major', default='')
    p.add_argument('--date', default='')
    p.set_defaults(func=cmd_template)

    args = ap.parse_args(argv)
    return args.func(args)


if __name__ == '__main__':
    raise SystemExit(main())
